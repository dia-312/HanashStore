"""Excel (.xlsx) and CSV product import (SPEC section 4, F2). Flagged for line-by-line review:
file upload, and it creates stock.

The client's real Excel file has not arrived yet (SPEC section 16). The columns below are the ones
the SPEC names (product, category, color, size, price, opening quantity), recognized by Arabic or
English header text. When the real file comes, only COLUMN_ALIASES (and maybe a header row offset)
should need to change.
"""

import csv
import io
import uuid
import zipfile
from collections.abc import Iterator
from dataclasses import dataclass, field
from typing import Any

from openpyxl import load_workbook
from sqlalchemy import select
from sqlalchemy.orm import Session

from app.audit import audit
from app.changeseq import take_seq
from app.errors import ApiError
from app.ids import uuid7
from app.models import Category, Product, ProductVariant
from app.money import MoneyError, parse_agorot, parse_quantity
from app.normalize import normalize_arabic
from app.schemas import PreviewRow, RowIssue
from app.services import catalog

MAX_BYTES = 5 * 1024 * 1024
MAX_ROWS = 5_000
MAX_SCAN_ROWS = 20_000  # raw rows looked at, so a sheet of a million blank rows can't stall us
MAX_UNZIPPED_BYTES = 100 * 1024 * 1024  # an .xlsx is a zip; refuse zip bombs
MAX_REPORTED_ERRORS = 200
PREVIEW_ROWS = 20

REQUIRED = ("product", "category", "price")
OPTIONAL = ("color", "size", "qty")

COLUMN_ALIASES: dict[str, tuple[str, ...]] = {
    "product": (
        "المنتج",
        "اسم المنتج",
        "الصنف",
        "اسم الصنف",
        "المادة",
        "اسم المادة",
        "product",
        "product name",
        "item",
        "name",
    ),
    "category": ("الفئة", "التصنيف", "القسم", "category"),
    "color": ("اللون", "color", "colour"),
    "size": ("الحجم", "المقاس", "size"),
    "price": ("السعر", "سعر البيع", "سعر", "price", "selling price"),
    "qty": (
        "الكمية",
        "الكمية الافتتاحية",
        "الرصيد",
        "المخزون",
        "الكمية المتوفرة",
        "quantity",
        "qty",
        "stock",
        "opening quantity",
    ),
}
_HEADER_TO_FIELD = {
    normalize_arabic(alias): name for name, aliases in COLUMN_ALIASES.items() for alias in aliases
}


@dataclass
class ParsedSheet:
    headers: list[str]
    rows: list[tuple[int, list[Any]]]  # (spreadsheet row number, cells)


@dataclass
class Report:
    rows: int = 0
    mapped: dict[str, str] = field(default_factory=dict)
    ignored: list[str] = field(default_factory=list)
    errors: list[RowIssue] = field(default_factory=list)
    errors_total: int = 0
    new_categories: int = 0
    new_products: int = 0
    new_variants: int = 0
    skipped_existing: int = 0
    opening_units: int = 0
    preview: list[PreviewRow] = field(default_factory=list)

    def error(self, row: int, field_name: str, code: str, detail: str | None = None) -> None:
        self.errors_total += 1
        if len(self.errors) < MAX_REPORTED_ERRORS:
            self.errors.append(RowIssue(row=row, field=field_name, code=code, detail=detail))


@dataclass
class PlannedVariant:
    row: int
    product_key: tuple[str, str]
    product_name: str
    category_name: str
    color: str | None
    size: str | None
    price: int
    qty: int


# ---------- reading the file ----------


def _cell_text(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))  # a size typed as 25 comes back as 25.0
    return str(value).strip()


def _read_xlsx(data: bytes) -> Iterator[list[Any]]:
    try:
        with zipfile.ZipFile(io.BytesIO(data)) as archive:
            if sum(i.file_size for i in archive.infolist()) > MAX_UNZIPPED_BYTES:
                raise ApiError(422, "IMPORT_TOO_LARGE", "The spreadsheet expands to too much data")
        workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        sheet = workbook.worksheets[0]
        for index, row in enumerate(sheet.iter_rows(values_only=True)):
            if index >= MAX_SCAN_ROWS:
                raise ApiError(422, "IMPORT_TOO_MANY_ROWS", f"More than {MAX_SCAN_ROWS} rows")
            yield list(row)
    except ApiError:
        raise
    except Exception as exc:  # corrupt zip, missing parts, bad XML: whatever openpyxl throws
        raise ApiError(422, "IMPORT_UNREADABLE", "The .xlsx file could not be read") from exc


def _read_csv(data: bytes) -> Iterator[list[Any]]:
    try:
        text = data.decode("utf-8-sig")
    except UnicodeDecodeError:
        text = data.decode("cp1256", errors="replace")  # Arabic Windows / old Excel exports
    first_line = text.split("\n", 1)[0]
    delimiter = max(",;\t", key=first_line.count) if any(d in first_line for d in ",;\t") else ","
    for index, row in enumerate(csv.reader(io.StringIO(text), delimiter=delimiter)):
        if index >= MAX_SCAN_ROWS:
            raise ApiError(422, "IMPORT_TOO_MANY_ROWS", f"More than {MAX_SCAN_ROWS} rows")
        yield list(row)


def parse_file(data: bytes) -> ParsedSheet:
    if len(data) > MAX_BYTES:
        raise ApiError(413, "IMPORT_TOO_LARGE", "The file is larger than 5 MB")
    if data[:4] == b"\xd0\xcf\x11\xe0":
        raise ApiError(422, "IMPORT_OLD_EXCEL", "Old .xls files are not supported; save as .xlsx")
    lines = _read_xlsx(data) if data[:4] == b"PK\x03\x04" else _read_csv(data)
    header: list[str] | None = None
    rows: list[tuple[int, list[Any]]] = []
    for number, cells in enumerate(lines, start=1):
        if all(_cell_text(c) == "" for c in cells):
            continue
        if header is None:
            header = [_cell_text(c) for c in cells]
            continue
        rows.append((number, cells))
        if len(rows) > MAX_ROWS:
            raise ApiError(422, "IMPORT_TOO_MANY_ROWS", f"More than {MAX_ROWS} product rows")
    if header is None:
        raise ApiError(422, "IMPORT_EMPTY", "The file has no rows")
    return ParsedSheet(headers=header, rows=rows)


# ---------- checking the rows ----------


def _map_columns(sheet: ParsedSheet, report: Report) -> dict[str, int]:
    columns: dict[str, int] = {}
    for index, header in enumerate(sheet.headers):
        name = _HEADER_TO_FIELD.get(normalize_arabic(header))
        if name is None or name in columns:
            if header:
                report.ignored.append(header)
            continue
        columns[name] = index
        report.mapped[name] = header
    for name in REQUIRED:
        if name not in columns:
            report.error(1, name, "MISSING_COLUMN")
    return columns


def _cell(cells: list[Any], columns: dict[str, int], name: str) -> Any:
    index = columns.get(name)
    return cells[index] if index is not None and index < len(cells) else None


def _check_rows(
    db: Session, sheet: ParsedSheet, columns: dict[str, int], report: Report
) -> list[PlannedVariant]:
    categories = {normalize_arabic(c.name): c for c in db.scalars(select(Category))}
    products = {(normalize_arabic(p.name), p.category_id): p for p in db.scalars(select(Product))}
    existing_variants = {
        (v.product_id, normalize_arabic(v.color or ""), normalize_arabic(v.size or ""))
        for v in db.scalars(select(ProductVariant))
    }
    planned: list[PlannedVariant] = []
    new_categories: set[str] = set()
    new_products: set[tuple[str, str]] = set()
    seen: dict[tuple[tuple[str, str], str, str], int] = {}
    report.rows = len(sheet.rows)

    for number, cells in sheet.rows:
        before = report.errors_total
        product_name = _cell_text(_cell(cells, columns, "product"))
        category_name = _cell_text(_cell(cells, columns, "category"))
        color = catalog.clean(_cell_text(_cell(cells, columns, "color")))
        size = catalog.clean(_cell_text(_cell(cells, columns, "size")))
        for name, value, limit in (
            ("product", product_name, 200),
            ("category", category_name, 200),
            ("color", color or "", 100),
            ("size", size or "", 100),
        ):
            if name in ("product", "category") and not value:
                report.error(number, name, "REQUIRED")
            elif len(value) > limit:
                report.error(number, name, "TOO_LONG", str(limit))
        price = qty = 0
        try:
            raw_price = _cell(cells, columns, "price")
            if raw_price is None or _cell_text(raw_price) == "":
                report.error(number, "price", "REQUIRED")
            else:
                price = parse_agorot(
                    raw_price if not isinstance(raw_price, str) else raw_price.strip()
                )
        except (MoneyError, ValueError, TypeError):
            report.error(number, "price", "BAD_PRICE", _cell_text(_cell(cells, columns, "price")))
        raw_qty = _cell(cells, columns, "qty")
        if raw_qty is not None and _cell_text(raw_qty) != "":
            try:
                qty = parse_quantity(raw_qty)
            except (MoneyError, ValueError, TypeError):
                report.error(number, "qty", "BAD_QTY", _cell_text(raw_qty))
        if report.errors_total > before:
            continue

        category = categories.get(normalize_arabic(category_name))
        if category is not None and not category.is_active:
            report.error(number, "category", "CATEGORY_INACTIVE", category_name)
            continue
        product_key = (normalize_arabic(product_name), normalize_arabic(category_name))
        existing_product = products.get((product_key[0], category.id)) if category else None
        if existing_product is not None and not existing_product.is_active:
            report.error(number, "product", "PRODUCT_INACTIVE", product_name)
            continue

        variant_key = (product_key, normalize_arabic(color or ""), normalize_arabic(size or ""))
        if variant_key in seen:
            report.error(number, "color", "DUPLICATE_IN_FILE", str(seen[variant_key]))
            continue
        seen[variant_key] = number
        if (
            existing_product is not None
            and (
                existing_product.id,
                variant_key[1],
                variant_key[2],
            )
            in existing_variants
        ):
            report.skipped_existing += 1  # never overwrite the price or stock of a known variant
            continue
        if category is None:
            new_categories.add(product_key[1])
        if existing_product is None:
            new_products.add(product_key)
        planned.append(
            PlannedVariant(
                number, product_key, product_name, category_name, color, size, price, qty
            )
        )
    report.new_categories = len(new_categories)
    report.new_products = len(new_products)
    return planned


# ---------- applying ----------


def _apply(
    db: Session, planned: list[PlannedVariant], report: Report, device_id: uuid.UUID | None
) -> None:
    seq = take_seq(db)
    categories = {normalize_arabic(c.name): c for c in db.scalars(select(Category))}
    products: dict[tuple[str, str], Product] = {}
    for p in db.scalars(select(Product)):
        category = db.get(Category, p.category_id)
        if category is not None:
            products[(normalize_arabic(p.name), normalize_arabic(category.name))] = p
    touched: dict[uuid.UUID, Product] = {}

    for item in planned:
        category_key = item.product_key[1]
        category = categories.get(category_key)
        if category is None:
            category = catalog.create_category(db, item.category_name, seq, device_id)
            categories[category_key] = category
        product = products.get(item.product_key)
        if product is None:
            product = catalog.create_product(
                db,
                name=item.product_name,
                category=category,
                description=None,
                seq=seq,
                device_id=device_id,
            )
            products[item.product_key] = product
        catalog.create_variant(
            db,
            product=product,
            color=item.color,
            size=item.size,
            price=item.price,
            min_stock=0,
            opening_qty=item.qty,
            seq=seq,
            device_id=device_id,
            refresh=False,
        )
        touched[product.id] = product
    for product in touched.values():
        catalog.refresh_search_text(db, product, seq)
    audit(
        db,
        "import",
        uuid7(),
        "products",
        after={
            "categories": report.new_categories,
            "products": report.new_products,
            "variants": report.new_variants,
            "skipped_existing": report.skipped_existing,
            "opening_units": report.opening_units,
        },
        device_id=device_id,
    )


def run_import(
    db: Session, data: bytes, *, dry_run: bool, device_id: uuid.UUID | None
) -> tuple[Report, bool]:
    """Returns (report, applied). Nothing is written unless there are no errors and not dry_run."""
    sheet = parse_file(data)
    report = Report()
    columns = _map_columns(sheet, report)
    planned = _check_rows(db, sheet, columns, report) if not report.errors_total else []

    report.new_variants = len(planned)
    report.opening_units = sum(p.qty for p in planned)
    report.preview = [
        PreviewRow(
            row=p.row,
            product=p.product_name,
            category=p.category_name,
            color=p.color,
            size=p.size,
            price=p.price,
            qty=p.qty,
        )
        for p in planned[:PREVIEW_ROWS]
    ]
    if dry_run or report.errors_total:
        return report, False
    _apply(db, planned, report, device_id)
    return report, True
